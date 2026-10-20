import json,openpyxl
from openpyxl.styles import Font
c=json.load(open('demo/game_config.json'));r=json.load(open('math/sim_results.json'))
wb=openpyxl.Workbook();ws=wb.active;ws.title="Shrnutí"
for k,v in r.items():
    if not isinstance(v,dict): ws.append([k,v])
ws.append([]);ws.append(["Jackpot zásahy"])
for k,v in r["jackpotHits"].items(): ws.append([k,v])
ws.append([]);ws.append(["Rozdělení výher (násobky celk. sázky)","počet"])
for k,v in r["winDist"].items(): ws.append([k,v])
p=wb.create_sheet("Výherní tabulka");p.append(["Symbol","Typ","3x","4x","5x","(násobky sázky na linii)"])
for s in c["symbols"]:
    if s["pays"]: p.append([s["name"],s["tier"]]+s["pays"][3:])
p.append(["Mapa pokladu (scatter)","celk. sázka"]+c["scatter"]["pays"][3:]+["3+ = 10 volných otáček"])
p.append(["Vykopej poklad ceny (x celk. sázka)"]+c["pickBonus"]["prizes"]+["lebek: %d"%c["pickBonus"]["skulls"]])
f=wb.create_sheet("Funkce")
k=c["krakenAttack"];f.append(["Útok krakena","pravděpodobnost/otáčka",k["probability"],"počet WILD",str(k["wildCounts"]),"válce",str(k["reels"])])
t=c["tidalWave"];f.append(["Přílivová vlna","pravděpodobnost po nevýherní otáčce",t["probability"],t["respin"]])
j=c["jackpot"];f.append(["Jackpot Kapitánova truhla","spuštění / otáčka při min. sázce",j["triggerProbPerMinBet"],"min. sázka",j["minBet"]])
f.append(["Úroveň","Výchozí (x min. sázka)","Příspěvek z sázky","Váha"])
for t in j["tiers"]: f.append([t["name"],t["seed"],t["contribution"],t["weight"]])
s=wb.create_sheet("Válce");s.append(["Pozice"]+["Válec %d"%(i+1) for i in range(5)])
for i in range(max(map(len,c["reels"]))): s.append([i]+[x[i] if i<len(x) else "" for x in c["reels"]])
s.append([]);s.append(["Délka"]+[len(x) for x in c["reels"]])
for sym in c["symbols"]: s.append([sym["id"]]+[x.count(sym["id"]) for x in c["reels"]])
l=wb.create_sheet("Výherní linie");l.append(["Linie"]+["V%d"%(i+1) for i in range(5)])
for i,x in enumerate(c["paylines"]): l.append([i+1]+x)
for w in wb: w["A1"].font=Font(bold=True)
wb.save("paytable_report.xlsx")
